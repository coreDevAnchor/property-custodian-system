"""Fix TOC numbering in the backup docx"""
import sys
import io
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

docx_path = r'D:\coredev\property-custodian-system\Property_Custodian_System_Documentation_backup.docx'
doc = Document(docx_path)

# TOC entries need to be renumbered
# Current -> New mapping (with Installation added as section 2)
toc_mapping = {
    "1.   Introduction & System Overview": "1.   Introduction & System Overview",  # unchanged
    "2.   Getting Started": "3.   Getting Started",
    "2.1   Logging In": "3.1   Logging In",
    "2.2   First-Time Password Change": "3.2   First-Time Password Change",
    "2.3   Understanding the Interface": "3.3   Understanding the Interface",
    "3.   Custodian Guide": "4.   Custodian Guide",
    "3.1   Dashboard": "4.1   Dashboard",
    "3.2   Asset Management": "4.2   Asset Management",
    "3.3   Managing Borrow Requests": "4.3   Managing Borrow Requests",
    "3.4   Managing Returns": "4.4   Managing Returns",
    "3.5   Employee Management": "4.5   Employee Management",
    "3.6   Custodian Management": "4.6   Custodian Management",
    "3.7   Audit Trail": "4.7   Audit Trail",
    "3.8   Reports & Analytics": "4.8   Reports & Analytics",
    "4.   Employee Guide": "5.   Employee Guide",
    "4.1   Dashboard": "5.1   Dashboard",
    "4.2   Browsing & Requesting Assets": "5.2   Browsing & Requesting Assets",
    "4.3   Managing Current Borrows": "5.3   Managing Current Borrows",
    "4.4   Borrow History": "5.4   Borrow History",
    "5.   Notifications": "6.   Notifications",
    "6.   Settings & Account": "7.   Settings & Account",
    "7.   Frequently Asked Questions": "8.   Frequently Asked Questions",
}

# Find TOC range (between "Table of Contents" heading and next heading)
toc_start = None
toc_end = None
for i, p in enumerate(doc.paragraphs):
    if p.style and p.style.name == 'Heading 1' and 'Table of Contents' in p.text:
        toc_start = i + 1
    if toc_start and i > toc_start and p.style and p.style.name == 'Heading 1':
        toc_end = i
        break

print(f"TOC range: {toc_start} to {toc_end}")

# First, insert "2. Installation" entry after Introduction
intro_idx = None
for i in range(toc_start, toc_end):
    p = doc.paragraphs[i]
    if '1.   Introduction' in p.text:
        intro_idx = i
        break

if intro_idx:
    # Insert installation entry after introduction
    from lxml import etree
    new_para = doc.add_paragraph("2.   Installation")
    new_para.style = doc.styles['Normal']
    # Move it after intro
    intro_elem = doc.paragraphs[intro_idx]._element
    intro_elem.addnext(new_para._element)
    print(f"Inserted '2.   Installation' after Introduction")
    toc_end += 1

# Now update existing entries
updated = 0
for i in range(toc_start, toc_end):
    p = doc.paragraphs[i]
    text = p.text.strip()
    if text in toc_mapping and toc_mapping[text] != text:
        # Clear and set new text
        for run in p.runs:
            run.text = ''
        if p.runs:
            p.runs[0].text = toc_mapping[text]
        else:
            p.add_run(toc_mapping[text])
        updated += 1
        print(f"Updated: '{text}' -> '{toc_mapping[text]}'")

print(f"\nUpdated {updated} TOC entries")

# Also renumber section headings in the body
heading_mapping = {
    "2. Getting Started": "3. Getting Started",
    "3. Custodian Guide": "4. Custodian Guide",
    "4. Employee Guide": "5. Employee Guide",
    "5. Notifications": "6. Notifications",
    "6. Settings & Account": "7. Settings & Account",
    "7. Frequently Asked Questions": "8. Frequently Asked Questions",
    "2.1 Logging In": "3.1 Logging In",
    "2.2 First-Time Password Change": "3.2 First-Time Password Change",
    "2.3 Understanding the Interface": "3.3 Understanding the Interface",
    "3.1 Dashboard": "4.1 Dashboard",
    "3.2 Asset Management": "4.2 Asset Management",
    "3.3 Managing Borrow Requests": "4.3 Managing Borrow Requests",
    "3.4 Managing Returns": "4.4 Managing Returns",
    "3.5 Employee Management": "4.5 Employee Management",
    "3.6 Custodian Management": "4.6 Custodian Management",
    "3.7 Audit Trail": "4.7 Audit Trail",
    "3.8 Reports & Analytics": "4.8 Reports & Analytics",
    "4.1 Employee Dashboard": "5.1 Employee Dashboard",
    "4.2 Browsing & Requesting Assets": "5.2 Browsing & Requesting Assets",
    "4.3 Managing Your Current Borrows": "5.3 Managing Your Current Borrows",
    "4.4 Borrow History": "5.4 Borrow History",
}

heading_updated = 0
for i, p in enumerate(doc.paragraphs):
    if p.style and p.style.name.startswith('Heading'):
        text = p.text.strip()
        if text in heading_mapping:
            for run in p.runs:
                run.text = ''
            if p.runs:
                p.runs[0].text = heading_mapping[text]
            else:
                p.add_run(heading_mapping[text])
            heading_updated += 1

print(f"Renumbered {heading_updated} headings")

doc.save(docx_path)
print(f"\nSaved to {docx_path}")
