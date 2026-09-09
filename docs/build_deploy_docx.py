#!/usr/bin/env python3
"""Build DEPLOY.docx, a reader-facing deployment guide for the Property Custodian app.

Run from repo root:  python docs/build_deploy_docx.py
Output: DEPLOY.docx (repo root). Requires python-docx:  pip install python-docx
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "DEPLOY.docx"

ENV_VARS = [
    ("APP_ENV", "production"),
    ("APP_DEBUG", "false"),
    ("APP_KEY", "___  Run  php artisan key:generate --show  and paste the output"),
    ("APP_URL", "https://<your-domain>  (must be the public URL users visit)"),
    ("TZ", "<your-timezone>"),
    ("DB_CONNECTION", "pgsql"),
    ("DB_HOST", "<pooled-host>.neon.tech"),
    ("DB_PORT", "5432"),
    ("DB_DATABASE", "<database-name>"),
    ("DB_USERNAME", "<database-user>"),
    ("DB_PASSWORD", "___  Copy from the Neon connection modal"),
    ("FILESYSTEM_DISK", "public  (the single disk the app uses; becomes R2 when R2_BUCKET is set)"),
    ("R2_ACCESS_KEY_ID", "___  Cloudflare R2 API token"),
    ("R2_SECRET_ACCESS_KEY", "___  Cloudflare R2 API token"),
    ("R2_BUCKET", "<your-bucket-name>"),
    ("R2_ENDPOINT", "https://<ACCOUNT_ID>.r2.cloudflarestorage.com"),
    ("SESSION_DRIVER", "database"),
    ("SESSION_LIFETIME", "120"),
    ("SESSION_SECURE_COOKIE", "true"),
    ("CACHE_STORE", "database"),
    ("QUEUE_CONNECTION", "database"),
    ("MAIL_FROM_NAME", '"Property Custodian"'),
]

STEPS = [
    (
        "Step 1 — Cloudflare R2 (stores photos)",
        [
            "Log in to Cloudflare and open R2 → Buckets → Create a bucket.",
            "Name it <your-bucket-name>. Region: APAC (Singapore).",
            "Keep the bucket PRIVATE (public access off). The app serves photos through its own proxy at /storage/{path}, so the frontend never talks to R2 directly.",
            "Open R2 → Manage R2 API Tokens → Create API token. Permission: Object Read & Write, scope: this bucket only.",
            "Copy and keep the Access Key ID and Secret Access Key — you will not see the secret again.",
            "Your endpoint is https://<your-account-id>.r2.cloudflarestorage.com (the Account ID is in your Cloudflare dashboard).",
        ],
    ),
    (
        "Step 2 — Neon (PostgreSQL database)",
        [
            "Create a free account at console.neon.tech → New project.",
            "Region: Singapore (ap-southeast-1). The free plan is enough.",
            "In the connection modal, copy the POOLED connection string (it uses port 5432 and a host ending in -pooler.neon.tech). You need the host name, database name and password for the env vars.",
        ],
    ),
    (
        "Step 3 — Render (the web service that runs the app)",
        [
            "Render is shown as one example host — the same steps work on any PHP 8.4 hosting.",
            "Create a free account at render.com → New → Web Service.",
            "Connect your Git repo and choose the branch to deploy.",
            "Runtime: Docker (Render builds the repo's Dockerfile automatically).",
            "Instance type: Free.",
            "Coat the service with the env vars in the table below, then Deploy.",
            "First boot automatically runs database migrations and seeds the built-in data (see docker/entrypoint.sh).",
        ],
    ),
    (
        "Step 4 — Pinger (keeps the free service awake)",
        [
            "Render's free tier sleeps after 15 minutes of no traffic and takes 30–60 s to wake up. A pinger keeps it awake.",
            "Use UptimeRobot (free) — HTTP(S) monitor, interval 10 min — or cron-job.org firing a GET every 10 minutes.",
            "URL to hit: https://<your-domain>/login",
            "This also keeps the 09:00 daily return-reminder mail running.",
        ],
    ),
    (
        "Step 5 — Verify after deploy",
        [
            "Open https://<your-domain>/login and register / sign in.",
            "Trigger a verification / OTP email and confirm it lands in the inbox.",
            "Upload a photo against an asset, then open its /storage/assets/... URL — the image should render.",
            "Open the custodian dashboard — the stats and recent activity render (no 500 or 502).",
        ],
    ),
]

VERIFY_TABLE = [
    ("Check", "What should happen"),
    ("Sign in", "Dashboard loads, no 500 or 502"),
    ("Send OTP email", "Verification mail arrives on any address"),
    ("Upload asset photo", "Photo is viewable via /storage/assets/..."),
    ("Custodian dashboard", "Stats + recent activity render"),
]

EMAIL_OPTIONS = [
    (
        "Option A — Brevo (recommended) — free ~300 emails/day",
        [
            "API-based, so it rides HTTPS 443 and works on any host including Render's free tier.",
            "Create a free account at brevo.com → Developers → API Keys, then Settings → Senders and verify the sender email.",
            "A 401 means a wrong API key; a 400 \"sender not verified\" means the sender step is pending.",
        ],
        "MAIL_MAILER=brevo\nBREVO_API_KEY=<your-brevo-api-key>\nMAIL_FROM_ADDRESS=<sender-email-verified-in-brevo>",
    ),
    (
        "Option B — Resend — free 3,000 emails/month",
        [
            "API-based. Requires the native package: composer require resend/resend-php.",
            "Free sandbox (sender onboarding@resend.dev) delivers ONLY to the inbox your Resend account uses — real recipients need a verified domain.",
        ],
        "MAIL_MAILER=resend\nRESEND_API_KEY=<your-resend-api-key>\nMAIL_FROM_ADDRESS=no-reply@<your-verified-domain>",
    ),
    (
        "Option C — Gmail SMTP — free (~500 emails/day with an existing Google account)",
        [
            "Turn on 2-Step Verification and create an App Password for the app.",
            "MAIL_FROM_ADDRESS must exactly equal MAIL_USERNAME, or Gmail rejects with 550 \"Sender address rejected\".",
            "Note: on Render's free tier outbound SMTP is frequently blocked/stalled (504) — the API options above are the reliable choice there. On VPS/shared hosting SMTP normally works.",
        ],
        "MAIL_MAILER=smtp\nMAIL_HOST=smtp.gmail.com\nMAIL_PORT=587\nMAIL_ENCRYPTION=tls\nMAIL_USERNAME=<your-gmail-address>\nMAIL_PASSWORD=<gmail-app-password>\nMAIL_FROM_ADDRESS=<your-gmail-address>",
    ),
]


def set_normal_style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def add_step(doc: Document, heading: str, bullets: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for i, bullet in enumerate(bullets, 1):
        p = doc.add_paragraph()
        p.add_run(f"{i}. ").bold = True
        p.add_run(bullet)


def add_env_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Variable"
    hdr[1].text = "Your value"
    for cell in hdr:
        for p in cell.paragraphs:
            p.add_run().font.bold = True
    for name, value in ENV_VARS:
        row = table.add_row().cells
        row[0].text = name
        row[1].text = value
    doc.add_paragraph("Secrets and account values are blanks — fill each one in from the account dashboards, then save and deploy.")


def main() -> None:
    doc = Document()
    set_normal_style(doc)

    title = doc.add_heading(0)
    title.add_run("Deploying the Property Custodian System")
    sub = doc.add_paragraph()
    sub.add_run(
        "A step-by-step guide to getting the app live, with a PostgreSQL database and "
        "Cloudflare R2 for photo storage. Allow about 30 minutes the first time. "
        "Nothing here is hosting-specific — the same steps (env vars + migrations) work "
        "on any PHP 8.4 host; Render's free tier is shown as one example."
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_heading("What you need", level=1)
    for item in [
        "Cloudflare account",
        "Neon account",
        "Render account (or any PHP 8.4 host)",
        "Email provider account (free tier): Brevo, Resend, or Gmail App Password — pick one",
        "UptimeRobot account (optional but recommended)",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    for heading, bullets in STEPS:
        add_step(doc, heading, bullets)

    doc.add_heading("Environment variables (paste into your host)", level=1)
    doc.add_paragraph(
        "In Render: web service → Environment. On any other host: the same values go in "
        ".env or the host's secret store. Fill every <placeholder> and blank with your own values."
    )
    add_code(
        doc,
        "APP_ENV=production\nAPP_DEBUG=false\nAPP_URL=https://<your-domain>\nTZ=<your-timezone>"
    )
    add_env_table(doc)

    doc.add_heading("Email — choose one provider (all free tiers)", level=1)
    for heading, bullets, code in EMAIL_OPTIONS:
        doc.add_heading(heading, level=2)
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")
        add_code(doc, code)

    doc.add_heading("Changed frontend code?", level=1)
    doc.add_paragraph(
        "The Docker image ships pre-built frontend assets from public/build (no Node in the container). "
        "Before pushing a deploy after frontend changes, rebuild and commit them:"
    )
    add_code(
        doc,
        "pnpm install --frozen-lockfile\npnpm run build\ngit add public/build\ngit commit -m \"build: assets\"\n"
    )

    doc.add_heading("Quick verify", level=1)
    for label, what in VERIFY_TABLE[1:]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(what)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()